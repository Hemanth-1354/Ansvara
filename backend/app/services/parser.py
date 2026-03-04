import io
import re
from typing import List
import PyPDF2
from docx import Document
import openpyxl


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text.strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def extract_text(file_bytes: bytes, filename: str) -> str:
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fname.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        return extract_text_from_xlsx(file_bytes)
    else:
        return extract_text_from_txt(file_bytes)


def parse_questions(text: str) -> List[str]:
    """Extract individual questions from questionnaire text."""
    questions = []

    # Try numbered list patterns: 1. 1) Q1. Q1:
    numbered = re.findall(
        r'(?:^|\n)\s*(?:Q\s*)?(\d+)[.):\s]+(.+?)(?=\n\s*(?:Q\s*)?\d+[.):\s]|\Z)',
        text,
        re.DOTALL | re.IGNORECASE
    )
    if numbered and len(numbered) >= 2:
        for _, q in numbered:
            q = q.strip()
            if len(q) > 10:
                questions.append(q)
        return questions

    # Try bullet/dash patterns
    bullets = re.findall(r'(?:^|\n)\s*[-•*]\s+(.+?)(?=\n\s*[-•*]|\Z)', text, re.DOTALL)
    if bullets and len(bullets) >= 2:
        for q in bullets:
            q = q.strip()
            if len(q) > 10:
                questions.append(q)
        return questions

    # Fallback: split by lines that end with ? or look like questions
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    for line in lines:
        if line.endswith('?') and len(line) > 15:
            questions.append(line)
        elif re.match(r'^(What|How|Does|Do|Is|Are|Can|Please|Describe|Explain|List)', line, re.I) and len(line) > 15:
            questions.append(line)

    if not questions:
        # Last resort: split by double newlines
        paragraphs = [p.strip() for p in text.split('\n\n') if p.strip() and len(p.strip()) > 20]
        questions = paragraphs[:15]

    return questions
